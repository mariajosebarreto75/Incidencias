"""
Servicio de Preoperacionales
Descarga InformePreoperacional.xlsx desde SharePoint (o ruta local)
y expone los datos procesados para el dashboard.
"""
import io
import logging
import os
from datetime import datetime
from threading import Lock
from typing import Optional

import requests

log = logging.getLogger(__name__)

# ── Columnas esperadas ────────────────────────────────────────────────────────
_COLUMNAS = [
    "Placa",
    "Tipo vehiculo",
    "Sede",
    "Contrato",
    "Fecha Carga",
    "Cruce Información",
    "Movimiento Vehículo",
    "Preoperacional",
    "Estado Preoperacional",
    "MES",
]

# Estados que cuentan como cumplimiento
_ESTADOS_OK = {
    "1. Preoperacional Confirmado",
    "2. Preoperacional Confirmado y sin GPS",
}
_ESTADO_SIN_MOVIMIENTO = "5. Vehículo sin movimiento"
_ESTADO_INCUMPLIMIENTO = "6. Vehículo con movimiento sin preoperacional"

# ── Cache en memoria + persistencia en disco ──────────────────────────────────
_cache: dict = {}
_cache_lock = Lock()

_DATA_DIR  = os.path.join(os.path.dirname(__file__), "..", "..", "data")
_CACHE_FILE = os.path.join(_DATA_DIR, "preop_cache.json")


def _guardar_cache_disco(registros: list[dict], meta: dict):
    """Serializa el cache a disco para que todos los workers puedan leerlo."""
    try:
        os.makedirs(_DATA_DIR, exist_ok=True)
        payload = {
            "meta": meta,
            "datos": [
                {**r, "fecha": r["fecha"].isoformat() if r.get("fecha") else None}
                for r in registros
            ],
        }
        tmp = _CACHE_FILE + ".tmp"
        import json as _json
        with open(tmp, "w", encoding="utf-8") as f:
            _json.dump(payload, f, ensure_ascii=False)
        os.replace(tmp, _CACHE_FILE)
    except Exception as e:
        log.warning("[Preop] No se pudo guardar cache en disco: %s", e)


def _cargar_cache_disco() -> tuple[list[dict], Optional[dict]]:
    """Carga el cache desde disco si existe."""
    try:
        import json as _json
        from datetime import date as _date
        with open(_CACHE_FILE, encoding="utf-8") as f:
            payload = _json.load(f)
        registros = []
        for r in payload.get("datos", []):
            fecha_str = r.get("fecha") or ""
            try:
                fecha = _date.fromisoformat(fecha_str) if fecha_str else None
            except Exception:
                fecha = None
            r["fecha"] = fecha
            registros.append(r)
        return registros, payload.get("meta")
    except Exception:
        return [], None


def _url() -> str:
    url = os.getenv("PREOP_SHAREPOINT_URL", "").strip()
    if not url:
        raise RuntimeError("Variable PREOP_SHAREPOINT_URL no configurada.")
    return url


def _descargar(url: str) -> bytes:
    """Descarga el archivo. Soporta URLs SharePoint y rutas locales."""
    if url.startswith(("http://", "https://")):
        # SharePoint: agregar &download=1 si no está presente
        if "download=" not in url:
            sep = "&" if "?" in url else "?"
            url = url + sep + "download=1"
        resp = requests.get(url, timeout=60)
        resp.raise_for_status()
        ct = resp.headers.get("Content-Type", "")
        if "html" in ct and b"<!DOCTYPE" in resp.content[:200]:
            raise RuntimeError(
                "SharePoint devolvió HTML en lugar del archivo. "
                "Verifica que el link sea de descarga directa y que el archivo sea público."
            )
        return resp.content
    else:
        # Ruta local (desarrollo / fallback)
        with open(url, "rb") as f:
            return f.read()


def _parsear(contenido: bytes) -> list[dict]:
    """Convierte el contenido binario del Excel en lista de dicts."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:
        raise RuntimeError(f"No se pudo abrir el archivo Excel: {e}")

    if "Ultimos9Dias" not in wb.sheetnames:
        raise RuntimeError(
            f"Hoja 'Ultimos9Dias' no encontrada. Hojas disponibles: {wb.sheetnames}"
        )

    ws = wb["Ultimos9Dias"]
    encabezados_raw = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    encabezados = [str(h).strip() if h is not None else "" for h in encabezados_raw]

    # Validar columnas críticas
    faltantes = [c for c in _COLUMNAS if c not in encabezados]
    if faltantes:
        log.warning("[Preop] Columnas no encontradas: %s", faltantes)

    # Mapeo índice por nombre
    idx = {h: i for i, h in enumerate(encabezados)}

    def _get(row, col):
        i = idx.get(col)
        return row[i] if i is not None and i < len(row) else None

    registros = []
    errores_fecha = 0

    for fila in ws.iter_rows(min_row=2, values_only=True):
        # Ignorar filas completamente vacías
        if all(v is None for v in fila):
            continue

        placa    = str(_get(fila, "Placa") or "").strip()
        tipo     = str(_get(fila, "Tipo vehiculo") or "").strip()
        sede     = str(_get(fila, "Sede") or "").strip()
        contrato = str(_get(fila, "Contrato") or "").strip()
        estado   = str(_get(fila, "Estado Preoperacional") or "").strip()
        preop    = str(_get(fila, "Preoperacional") or "").strip()
        mes      = str(_get(fila, "MES") or "").strip()

        # Movimiento: puede ser float o texto
        mov_raw  = _get(fila, "Movimiento Vehículo")
        try:
            movimiento = float(mov_raw) if mov_raw not in (None, "", "Sin informacion") else None
        except (TypeError, ValueError):
            movimiento = None

        # Fecha
        fecha_raw = _get(fila, "Fecha Carga")
        if isinstance(fecha_raw, datetime):
            fecha = fecha_raw.date()
        else:
            try:
                fecha = datetime.strptime(str(fecha_raw), "%Y-%m-%d").date() if fecha_raw else None
            except Exception:
                fecha = None
                errores_fecha += 1

        if not placa or not estado:
            continue

        registros.append({
            "placa":      placa,
            "tipo":       tipo,
            "sede":       sede,
            "contrato":   contrato,
            "fecha":      fecha,
            "fecha_str":  str(fecha) if fecha else "",
            "movimiento": movimiento,
            "preop":      preop,
            "estado":     estado,
            "mes":        mes,
            "cumple":     estado in _ESTADOS_OK,
            "sin_mov":    estado == _ESTADO_SIN_MOVIMIENTO,
            "incumple":   estado == _ESTADO_INCUMPLIMIENTO,
        })

    if errores_fecha:
        log.warning("[Preop] %d fechas inválidas ignoradas.", errores_fecha)

    log.info("[Preop] %d registros cargados.", len(registros))
    return registros


def sincronizar() -> dict:
    """Descarga y procesa el Excel. Actualiza el cache. Devuelve resumen."""
    url = _url()
    log.info("[Preop] Iniciando sincronización desde %s…", url[:60])
    try:
        contenido = _descargar(url)
        registros = _parsear(contenido)
        resultado = {
            "ok": True,
            "registros": len(registros),
            "actualizado": datetime.utcnow().isoformat(),
            "error": None,
        }
        with _cache_lock:
            _cache["datos"] = registros
            _cache["meta"]  = resultado
        _guardar_cache_disco(registros, resultado)
        log.info("[Preop] Sincronización OK: %d registros.", len(registros))
        return resultado
    except Exception as e:
        log.error("[Preop] Error en sincronización: %s", e)
        resultado = {
            "ok": False,
            "registros": 0,
            "actualizado": datetime.utcnow().isoformat(),
            "error": str(e),
        }
        with _cache_lock:
            _cache.setdefault("datos", [])
            _cache["meta"] = resultado
        return resultado


def obtener_datos() -> tuple[list[dict], Optional[dict]]:
    """Devuelve (registros, meta). Lee de disco si el cache en memoria está vacío."""
    with _cache_lock:
        if not _cache.get("datos"):
            datos_disco, meta_disco = _cargar_cache_disco()
            if datos_disco:
                _cache["datos"] = datos_disco
                _cache["meta"]  = meta_disco
        return _cache.get("datos", []), _cache.get("meta")


def _kpis(registros: list[dict], filtros: dict) -> dict:
    """Calcula KPIs sobre los registros ya filtrados."""
    total = len(registros)
    if total == 0:
        return {"total": 0, "cumple": 0, "incumple": 0, "sin_mov": 0,
                "pct_cumplimiento": 0, "pct_incumplimiento": 0}

    cumple   = sum(1 for r in registros if r["cumple"])
    incumple = sum(1 for r in registros if r["incumple"])
    sin_mov  = sum(1 for r in registros if r["sin_mov"])
    con_mov  = total - sin_mov  # vehículos que debían tener preoperacional

    pct_cumplimiento    = round(cumple / con_mov * 100, 1) if con_mov else 0
    pct_incumplimiento  = round(incumple / con_mov * 100, 1) if con_mov else 0

    return {
        "total":              total,
        "cumple":             cumple,
        "incumple":           incumple,
        "sin_mov":            sin_mov,
        "con_movimiento":     con_mov,
        "pct_cumplimiento":   pct_cumplimiento,
        "pct_incumplimiento": pct_incumplimiento,
    }


def _filtrar(registros: list[dict], filtros: dict) -> list[dict]:
    sede     = filtros.get("sede")
    contrato = filtros.get("contrato")
    fecha    = filtros.get("fecha")
    tipo     = filtros.get("tipo")

    resultado = registros
    if sede:
        resultado = [r for r in resultado if r["sede"] == sede]
    if contrato:
        resultado = [r for r in resultado if r["contrato"] == contrato]
    if fecha:
        resultado = [r for r in resultado if r["fecha_str"] == fecha]
    if tipo:
        resultado = [r for r in resultado if r["tipo"] == tipo]
    return resultado


def datos_dashboard(filtros: dict | None = None) -> dict:
    """
    Retorna todo lo necesario para el dashboard:
    kpis, serie temporal, distribución por estado, top incumplimientos,
    listas de filtros disponibles.
    """
    from collections import Counter, defaultdict

    registros, meta = obtener_datos()
    filtros = filtros or {}
    filtrados = _filtrar(registros, filtros)

    kpis = _kpis(filtrados, filtros)

    # Serie por fecha
    serie: dict = defaultdict(lambda: {"total": 0, "cumple": 0, "incumple": 0})
    for r in filtrados:
        d = r["fecha_str"] or "sin fecha"
        serie[d]["total"]    += 1
        if r["cumple"]:   serie[d]["cumple"]   += 1
        if r["incumple"]: serie[d]["incumple"] += 1
    serie_lista = [{"fecha": k, **v} for k, v in sorted(serie.items())]

    # Distribución por estado
    estados_cnt = Counter(r["estado"] for r in filtrados)
    estados_lista = [{"estado": k, "total": v}
                     for k, v in sorted(estados_cnt.items(), key=lambda x: x[0])]

    # Top contratos con incumplimiento
    inc_contrato = Counter(r["contrato"] for r in filtrados if r["incumple"])
    top_incumplimiento = [{"contrato": k, "incumple": v}
                          for k, v in inc_contrato.most_common(10)]

    # Cumplimiento por sede
    sedes_data: dict = defaultdict(lambda: {"total": 0, "cumple": 0})
    for r in filtrados:
        s = r["sede"] or "Sin sede"
        sedes_data[s]["total"] += 1
        if r["cumple"]:
            sedes_data[s]["cumple"] += 1
    sedes_lista = []
    for sede, d in sorted(sedes_data.items()):
        con_mov = d["total"] - sum(1 for r in filtrados if r["sede"] == sede and r["sin_mov"])
        pct = round(d["cumple"] / con_mov * 100, 1) if con_mov else 0
        sedes_lista.append({"sede": sede, "total": d["total"],
                             "cumple": d["cumple"], "pct": pct})

    # Opciones de filtros
    opciones = {
        "sedes":     sorted({r["sede"]     for r in registros if r["sede"]}),
        "contratos": sorted({r["contrato"] for r in registros if r["contrato"]}),
        "fechas":    sorted({r["fecha_str"] for r in registros if r["fecha_str"]}),
        "tipos":     sorted({r["tipo"]     for r in registros if r["tipo"]}),
    }

    return {
        "kpis":              kpis,
        "serie_temporal":    serie_lista,
        "estados":           estados_lista,
        "top_incumplimiento": top_incumplimiento,
        "sedes":             sedes_lista,
        "opciones":          opciones,
        "meta":              meta,
        "filtros_activos":   filtros,
    }
