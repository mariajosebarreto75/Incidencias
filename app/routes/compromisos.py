import os
import uuid
from datetime import date, datetime

from flask import (Blueprint, render_template, request, redirect, url_for,
                   flash, jsonify, send_from_directory, current_app, abort)
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models.compromiso import Compromiso, HistorialReprogramacion
from app.models.contrato import Contrato
from app.models.user_contrato import UserContrato

compromisos_bp = Blueprint("compromisos", __name__, url_prefix="/compromisos")


def _es_neo():
    """Retorna True si el usuario tiene permisos completos (neo o admin)."""
    return current_user.rol.lower() in ("neo", "admin")

ALLOWED_EXT = {"png", "jpg", "jpeg", "webp", "pdf"}
UPLOAD_FOLDER_NAME = "evidencias_compromisos"


def _upload_dir():
    base = os.path.join(current_app.root_path, "uploads", UPLOAD_FOLDER_NAME)
    os.makedirs(base, exist_ok=True)
    return base


def _allowed(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXT


def _contratos_usuario():
    """Contratos que el usuario actual puede ver."""
    if current_user.rol in ("admin", "neo"):
        return Contrato.query.filter_by(activo=True).order_by(Contrato.contrato).all()
    # coordinador: solo sus contratos asignados
    asignados = (UserContrato.query
                 .filter_by(user_id=current_user.id)
                 .with_entities(UserContrato.contrato).all())
    nombres = [a.contrato for a in asignados]
    return (Contrato.query
            .filter(Contrato.contrato.in_(nombres), Contrato.activo == True)
            .order_by(Contrato.contrato).all())


def _puede_ver_contrato(contrato_id):
    contratos = _contratos_usuario()
    return any(c.id == contrato_id for c in contratos)


# ── LISTADO ──────────────────────────────────────────────────────────────────

@compromisos_bp.route("/")
@login_required
def index():
    contratos = _contratos_usuario()
    contrato_id   = request.args.get("contrato_id", type=int)
    atrasado_filtro = request.args.get("atrasado", "")
    q_resp        = request.args.get("q_resp", "").strip()
    fecha_desde   = request.args.get("fecha_desde", "")
    fecha_hasta   = request.args.get("fecha_hasta", "")

    # checkboxes de estado (ausente = todos marcados)
    est_pend = request.args.get("est_pend")
    est_rep  = request.args.get("est_rep")
    est_cerr = request.args.get("est_cerr")
    # si ninguno fue enviado (GET limpio) → todos activos
    estados_sel = []
    if est_pend or est_rep or est_cerr:
        if est_pend: estados_sel.append("Pendiente")
        if est_rep:  estados_sel.append("Reprogramado")
        if est_cerr: estados_sel.append("Cerrado")
    # si se enviaron pero ninguno marcado → igual todos (evita lista vacía)
    if not estados_sel:
        estados_sel = ["Pendiente", "Reprogramado", "Cerrado"]

    q = Compromiso.query.join(Contrato).filter(
        Contrato.id.in_([c.id for c in contratos])
    )

    if contrato_id:
        q = q.filter(Compromiso.contrato_id == contrato_id)

    q = q.filter(Compromiso.estado.in_(estados_sel))

    if q_resp:
        q = q.filter(Compromiso.responsable.ilike(f"%{q_resp}%"))
    if fecha_desde:
        try:
            q = q.filter(Compromiso.fecha_entrega >= date.fromisoformat(fecha_desde))
        except ValueError:
            pass
    if fecha_hasta:
        try:
            q = q.filter(Compromiso.fecha_entrega <= date.fromisoformat(fecha_hasta))
        except ValueError:
            pass

    compromisos = q.order_by(Compromiso.fecha_entrega.asc(), Compromiso.id.desc()).all()

    if atrasado_filtro == "1":
        compromisos = [c for c in compromisos if c.atrasado]
    elif atrasado_filtro == "0":
        compromisos = [c for c in compromisos if not c.atrasado]

    total        = len(compromisos)
    pendientes   = sum(1 for c in compromisos if c.estado == "Pendiente")
    reprogramados = sum(1 for c in compromisos if c.estado == "Reprogramado")
    cerrados     = sum(1 for c in compromisos if c.estado == "Cerrado")
    atrasados    = sum(1 for c in compromisos if c.atrasado)

    return render_template(
        "compromisos/index.html",
        compromisos=compromisos,
        contratos=contratos,
        contrato_id_sel=contrato_id,
        atrasado_sel=atrasado_filtro,
        q_resp_sel=q_resp,
        fecha_desde_sel=fecha_desde,
        fecha_hasta_sel=fecha_hasta,
        est_pend_sel=est_pend,
        est_rep_sel=est_rep,
        est_cerr_sel=est_cerr,
        es_neo=_es_neo(),
        kpi=dict(total=total, pendientes=pendientes,
                 reprogramados=reprogramados, cerrados=cerrados, atrasados=atrasados),
    )


# ── CREAR ─────────────────────────────────────────────────────────────────────

@compromisos_bp.route("/nuevo", methods=["GET", "POST"])
@login_required
def nuevo():
    if not _es_neo():
        abort(403)
    contratos = _contratos_usuario()
    if request.method == "POST":
        contrato_id = request.form.get("contrato_id", type=int)
        responsable = request.form.get("responsable", "").strip()
        texto = request.form.get("compromiso", "").strip()
        fecha_entrega_str = request.form.get("fecha_entrega", "")
        obs = request.form.get("observacion_general", "").strip()

        errores = []
        if not contrato_id:
            errores.append("Selecciona un contrato.")
        if not responsable:
            errores.append("El responsable es obligatorio.")
        if not texto:
            errores.append("El texto del compromiso es obligatorio.")
        if not fecha_entrega_str:
            errores.append("La fecha de entrega es obligatoria.")
        else:
            try:
                fecha_entrega = date.fromisoformat(fecha_entrega_str)
            except ValueError:
                errores.append("Fecha inválida.")
                fecha_entrega = None

        if not _puede_ver_contrato(contrato_id):
            errores.append("No tienes acceso a ese contrato.")

        if errores:
            for e in errores:
                flash(e, "danger")
            return render_template("compromisos/form.html", contratos=contratos,
                                   form=request.form)

        comp = Compromiso(
            contrato_id=contrato_id,
            responsable=responsable,
            compromiso=texto,
            fecha_entrega=fecha_entrega,
            observacion_general=obs or None,
        )
        db.session.add(comp)
        db.session.commit()
        flash("Compromiso creado.", "success")
        return redirect(url_for("compromisos.index"))

    return render_template("compromisos/form.html", contratos=contratos, form={})


# ── DETALLE ───────────────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>")
@login_required
def detalle(comp_id):
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    return render_template("compromisos/detalle.html", comp=comp)


# ── REPROGRAMAR ───────────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/reprogramar", methods=["POST"])
@login_required
def reprogramar(comp_id):
    if not _es_neo():
        abort(403)
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    if comp.estado == "Cerrado":
        flash("No se puede reprogramar un compromiso cerrado.", "warning")
        return redirect(url_for("compromisos.index"))

    nueva_str = request.form.get("nueva_fecha", "")
    try:
        nueva = date.fromisoformat(nueva_str)
    except ValueError:
        flash("Fecha inválida.", "danger")
        return redirect(url_for("compromisos.index"))

    hist = HistorialReprogramacion(
        compromiso_id=comp.id,
        fecha_anterior=comp.fecha_entrega,
        nueva_fecha=nueva,
    )
    db.session.add(hist)
    comp.fecha_entrega = nueva
    comp.estado = "Reprogramado"
    comp.cantidad_reprogramaciones += 1
    db.session.commit()
    flash("Compromiso reprogramado.", "success")
    return redirect(url_for("compromisos.index"))


# ── CERRAR ────────────────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/cerrar", methods=["POST"])
@login_required
def cerrar(comp_id):
    if not _es_neo():
        abort(403)
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    if comp.estado == "Cerrado":
        flash("El compromiso ya está cerrado.", "warning")
        return redirect(url_for("compromisos.index"))

    fecha_real_str = request.form.get("fecha_entrega_real", "")
    obs_cierre = request.form.get("obs_cierre", "").strip()

    try:
        fecha_real = date.fromisoformat(fecha_real_str)
    except ValueError:
        flash("Fecha de cierre inválida.", "danger")
        return redirect(url_for("compromisos.index"))

    # Evidencia obligatoria para cerrar
    archivo = request.files.get("evidencia")
    tiene_evidencia = bool(comp.evidencia_path)
    if archivo and archivo.filename:
        if _allowed(archivo.filename):
            ext = archivo.filename.rsplit(".", 1)[1].lower()
            nombre_unico = f"{uuid.uuid4().hex}.{ext}"
            archivo.save(os.path.join(_upload_dir(), nombre_unico))
            comp.evidencia_path = nombre_unico
            comp.evidencia_nombre = secure_filename(archivo.filename)
            tiene_evidencia = True

    if not tiene_evidencia:
        flash("Para cerrar el compromiso primero debes subir una evidencia.", "warning")
        return redirect(url_for("compromisos.index"))

    if obs_cierre:
        sello = datetime.now().strftime("%Y-%m-%d %H:%M")
        actual = comp.observacion_general or ""
        comp.observacion_general = (f"{actual}\n\n[{sello}] {obs_cierre}" if actual
                                    else f"[{sello}] {obs_cierre}")

    comp.fecha_entrega_real = fecha_real
    comp.estado = "Cerrado"
    db.session.commit()
    flash("Compromiso cerrado.", "success")
    return redirect(url_for("compromisos.index"))


# ── SUBIR EVIDENCIA (standalone) ───────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/evidencia", methods=["POST"])
@login_required
def subir_evidencia(comp_id):
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    archivo = request.files.get("evidencia")
    if not archivo or not archivo.filename:
        flash("No se recibió archivo.", "danger")
        return redirect(url_for("compromisos.index"))
    if not _allowed(archivo.filename):
        flash("Tipo de archivo no permitido (jpg, png, webp, pdf).", "danger")
        return redirect(url_for("compromisos.index"))

    # Borrar evidencia anterior si existe
    if comp.evidencia_path:
        old = os.path.join(_upload_dir(), comp.evidencia_path)
        if os.path.exists(old):
            os.remove(old)

    ext = archivo.filename.rsplit(".", 1)[1].lower()
    nombre_unico = f"{uuid.uuid4().hex}.{ext}"
    archivo.save(os.path.join(_upload_dir(), nombre_unico))
    comp.evidencia_path = nombre_unico
    comp.evidencia_nombre = secure_filename(archivo.filename)
    db.session.commit()
    flash("Evidencia guardada.", "success")
    return redirect(url_for("compromisos.index"))


# ── VER EVIDENCIA ──────────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/evidencia/ver")
@login_required
def ver_evidencia(comp_id):
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    if not comp.evidencia_path:
        abort(404)
    return send_from_directory(_upload_dir(), comp.evidencia_path,
                               download_name=comp.evidencia_nombre or comp.evidencia_path)


# ── ELIMINAR EVIDENCIA ─────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/evidencia/eliminar", methods=["POST"])
@login_required
def eliminar_evidencia(comp_id):
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    if comp.evidencia_path:
        old = os.path.join(_upload_dir(), comp.evidencia_path)
        if os.path.exists(old):
            os.remove(old)
        comp.evidencia_path = None
        comp.evidencia_nombre = None
        db.session.commit()
    flash("Evidencia eliminada.", "success")
    return redirect(url_for("compromisos.index"))


# ── ELIMINAR COMPROMISO ────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/eliminar", methods=["POST"])
@login_required
def eliminar(comp_id):
    if not _es_neo():
        abort(403)
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    if comp.evidencia_path:
        old = os.path.join(_upload_dir(), comp.evidencia_path)
        if os.path.exists(old):
            os.remove(old)
    db.session.delete(comp)
    db.session.commit()
    flash("Compromiso eliminado.", "success")
    return redirect(url_for("compromisos.index"))


# ── OBSERVACIÓN (AJAX) ─────────────────────────────────────────────────────────

@compromisos_bp.route("/<int:comp_id>/observacion", methods=["POST"])
@login_required
def guardar_obs(comp_id):
    comp = Compromiso.query.get_or_404(comp_id)
    if not _puede_ver_contrato(comp.contrato_id):
        abort(403)
    texto = request.form.get("texto", "").strip()
    modo = request.form.get("modo", "append")
    if not texto:
        return jsonify(ok=False, error="Texto vacío"), 400

    if modo == "append":
        sello = datetime.now().strftime("%Y-%m-%d %H:%M")
        actual = comp.observacion_general or ""
        comp.observacion_general = (f"{actual}\n\n[{sello}] {texto}" if actual
                                    else f"[{sello}] {texto}")
    else:
        comp.observacion_general = texto
    db.session.commit()
    return jsonify(ok=True, obs=comp.observacion_general)


# ── CREAR CONTRATO RÁPIDO (AJAX) ───────────────────────────────────────────────

@compromisos_bp.route("/contratos/nuevo", methods=["POST"])
@login_required
def crear_contrato_rapido():
    if not _es_neo():
        return jsonify(ok=False, error="Sin permisos"), 403
    nombre = request.form.get("nombre", "").strip()
    if not nombre:
        return jsonify(ok=False, error="El nombre es obligatorio"), 400
    if Contrato.query.filter_by(contrato=nombre).first():
        return jsonify(ok=False, error="Ya existe un contrato con ese nombre"), 409
    c = Contrato(contrato=nombre, activo=True)
    db.session.add(c)
    db.session.commit()
    return jsonify(ok=True, id=c.id, nombre=c.contrato)


# ── ELIMINAR MASIVO ────────────────────────────────────────────────────────────

@compromisos_bp.route("/eliminar-masivo", methods=["POST"])
@login_required
def eliminar_masivo():
    if not _es_neo():
        abort(403)
    ids_raw = request.form.get("ids", "")
    try:
        ids = [int(x) for x in ids_raw.split(",") if x.strip().isdigit()]
    except ValueError:
        flash("IDs inválidos.", "danger")
        return redirect(url_for("compromisos.index"))

    eliminados = 0
    for comp_id in ids:
        comp = Compromiso.query.get(comp_id)
        if comp and _puede_ver_contrato(comp.contrato_id):
            if comp.evidencia_path:
                old = os.path.join(_upload_dir(), comp.evidencia_path)
                if os.path.exists(old):
                    os.remove(old)
            db.session.delete(comp)
            eliminados += 1

    db.session.commit()
    flash(f"{eliminados} compromiso(s) eliminado(s).", "success")
    return redirect(url_for("compromisos.index"))


# ── EXPORTAR EXCEL ─────────────────────────────────────────────────────────────

@compromisos_bp.route("/exportar-excel")
@login_required
def exportar_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from flask import Response
        import io
    except ImportError:
        flash("Instala openpyxl para exportar Excel.", "danger")
        return redirect(url_for("compromisos.index"))

    contratos = _contratos_usuario()
    contrato_id = request.args.get("contrato_id", type=int)

    q = Compromiso.query.join(Contrato).filter(
        Contrato.id.in_([c.id for c in contratos])
    )
    if contrato_id:
        q = q.filter(Compromiso.contrato_id == contrato_id)
    compromisos = q.order_by(Compromiso.fecha_entrega.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compromisos"

    cabeceras = ["ID", "Contrato", "Responsable", "Compromiso", "Creación",
                 "Entrega", "Estado", "Reprog.", "Cierre real", "Atrasado", "Observación"]
    hdr_fill = PatternFill("solid", fgColor="0D6E6E")
    hdr_font = Font(bold=True, color="FFFFFF")

    ws.append(cabeceras)
    for i, cell in enumerate(ws[1], 1):
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    for c in compromisos:
        ws.append([
            c.id,
            c.contrato.contrato,
            c.responsable,
            c.compromiso,
            c.fecha_creacion.strftime("%Y-%m-%d"),
            c.fecha_entrega.strftime("%Y-%m-%d"),
            c.estado,
            c.cantidad_reprogramaciones,
            c.fecha_entrega_real.strftime("%Y-%m-%d") if c.fecha_entrega_real else "",
            "Sí" if c.atrasado else "No",
            c.observacion_general or "",
        ])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=compromisos.xlsx"},
    )
